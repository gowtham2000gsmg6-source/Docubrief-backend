from pathlib import Path

from docx import Document

from app.services.extractors.base import ExtractedDocument, Extractor, TextSegment, combine_segments


class DocxExtractor(Extractor):
    def extract(self, path: Path) -> ExtractedDocument:
        document = Document(path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        tables: list[str] = []
        for table_index, table in enumerate(document.tables, start=1):
            rows = [
                " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
                for row in table.rows
            ]
            tables.append(f"Table {table_index}:\n" + "\n".join(rows))
        contents = "\n".join([*paragraphs, *tables]).strip()
        segment = TextSegment(marker="Document", text=contents)
        return ExtractedDocument(
            text=combine_segments([segment]),
            page_count=0,
            segments=[segment],
        )
