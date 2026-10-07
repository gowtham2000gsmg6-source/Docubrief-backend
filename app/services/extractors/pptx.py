from pathlib import Path

from pptx import Presentation

from app.services.extractors.base import ExtractedDocument, Extractor, TextSegment, combine_segments


class PptxExtractor(Extractor):
    def extract(self, path: Path) -> ExtractedDocument:
        presentation = Presentation(path)
        segments: list[TextSegment] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            blocks = [
                shape.text.strip()
                for shape in slide.shapes
                if shape.has_text_frame and shape.text.strip()
            ]
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except (AttributeError, KeyError):
                notes = ""
            if notes:
                blocks.append(f"Speaker notes:\n{notes}")
            segments.append(TextSegment(marker=f"Slide {slide_number}", text="\n".join(blocks)))
        return ExtractedDocument(
            text=combine_segments(segments),
            page_count=len(presentation.slides),
            segments=segments,
        )
