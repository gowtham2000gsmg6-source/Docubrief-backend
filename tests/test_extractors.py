from pathlib import Path

import fitz
import pytest
from docx import Document
from pptx import Presentation

from app.services.extractors import (
    DocxExtractor,
    MarkdownExtractor,
    PdfExtractor,
    PptxExtractor,
    TextFileExtractor,
    get_extractor,
)
from app.services.extractors.pdf import _layout_order


def test_txt_extractor_reads_utf8_and_marks_document(tmp_path: Path) -> None:
    source = tmp_path / "notes.txt"
    source.write_text("A readable note.\nSecond line.", encoding="utf-8")

    result = TextFileExtractor().extract(source)

    assert result.page_count == 1
    assert result.segments[0].marker == "Document"
    assert "[Document]\nA readable note.\nSecond line." == result.text


def test_txt_extractor_detects_utf16(tmp_path: Path) -> None:
    source = tmp_path / "notes.txt"
    source.write_bytes("Café résumé".encode("utf-16"))

    result = TextFileExtractor().extract(source)

    assert "Café résumé" in result.text


def test_markdown_extractor_preserves_markdown(tmp_path: Path) -> None:
    source = tmp_path / "readme.md"
    source.write_text("# Heading\n\n- first\n- second", encoding="utf-8")

    result = MarkdownExtractor().extract(source)

    assert result.page_count == 1
    assert result.segments[0].marker == "Markdown document"
    assert "# Heading" in result.text
    assert "- first\n- second" in result.text


def test_docx_extractor_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    source = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Quarterly report")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "$10m"
    document.save(source)

    result = DocxExtractor().extract(source)

    assert result.page_count == 0
    assert "[Document]" in result.text
    assert "Quarterly report" in result.text
    assert "Table 1:" in result.text
    assert "Metric | Value" in result.text
    assert "Revenue | $10m" in result.text


def test_pptx_extractor_includes_slide_text_and_speaker_notes(tmp_path: Path) -> None:
    source = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Launch plan"
    slide.placeholders[1].text = "First milestone"
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.text = "Remind the team about the launch date."
    presentation.save(source)

    result = PptxExtractor().extract(source)

    assert result.page_count == 1
    assert result.segments[0].marker == "Slide 1"
    assert "[Slide 1]" in result.text
    assert "Launch plan" in result.text
    assert "First milestone" in result.text
    assert "Speaker notes:" in result.text
    assert "Remind the team about the launch date." in result.text


def test_pdf_extractor_preserves_page_markers_and_count(tmp_path: Path) -> None:
    source = tmp_path / "report.pdf"
    with fitz.open() as document:
        for text in ("First page content", "Second page content"):
            page = document.new_page()
            page.insert_text((72, 72), text)
        document.save(source)

    result = PdfExtractor().extract(source)

    assert result.page_count == 2
    assert [segment.marker for segment in result.segments] == ["Page 1", "Page 2"]
    assert "[Page 1]" in result.text
    assert "First page content" in result.text
    assert "[Page 2]" in result.text
    assert "Second page content" in result.text


def test_pdf_extractor_does_not_require_ocr_for_short_selectable_text(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "short-text.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "Short text")
        document.save(source)

    def unexpected_ocr(*args: object, **kwargs: object) -> str:
        raise AssertionError("OCR should not run when selectable text was extracted.")

    monkeypatch.setattr("app.services.extractors.pdf.pytesseract.image_to_string", unexpected_ocr)

    result = PdfExtractor().extract(source)

    assert "Short text" in result.text


def test_pdf_extractor_uses_ocr_for_scanned_pages(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "scan.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(source)
    monkeypatch.setattr(
        "app.services.extractors.pdf.pytesseract.image_to_string",
        lambda image, lang: "OCR recognized document text",
    )

    result = PdfExtractor().extract(source)

    assert result.page_count == 1
    assert result.segments[0].text == "OCR recognized document text"
    assert "OCR recognized document text" in result.text


def test_pdf_layout_orders_columns_left_to_right() -> None:
    # PyMuPDF block tuples are (x0, y0, x1, y1, text, ...).
    blocks = [
        (250, 10, 350, 25, "right top", 0, 0),
        (10, 30, 110, 45, "left bottom", 0, 0),
        (10, 10, 110, 25, "left top", 0, 0),
        (250, 30, 350, 45, "right bottom", 0, 0),
    ]

    assert _layout_order(blocks) == [
        "left top",
        "left bottom",
        "right top",
        "right bottom",
    ]


@pytest.mark.parametrize(
    ("file_type", "extractor_type"),
    [
        ("pdf", PdfExtractor),
        ("docx", DocxExtractor),
        ("pptx", PptxExtractor),
        ("txt", TextFileExtractor),
        ("md", MarkdownExtractor),
    ],
)
def test_extractor_factory_maps_supported_types(file_type: str, extractor_type: type) -> None:
    assert isinstance(get_extractor(file_type), extractor_type)


def test_extractor_factory_rejects_unsupported_type() -> None:
    with pytest.raises(ValueError, match="Unsupported document type"):
        get_extractor("xlsx")
